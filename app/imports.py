import csv
import io
import json
import zipfile
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from .models import ImportErrorRow, ImportJob, ModelDict
from .schemas import LeadInput
from .services import create_lead, duplicate, serialize

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
FIELDS = {'full_name', 'phone', 'email', 'company_name', 'inquiry_type', 'priority'}


def parse_file(filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.csv', '.xlsx'}:
        raise HTTPException(422, 'Only CSV and XLSX files are supported')
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Maximum upload size is 5 MB')
    try:
        if suffix == '.csv':
            reader = csv.reader(io.StringIO(content.decode('utf-8-sig')))
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise HTTPException(422, 'Maximum 2,000 data rows per import')
        else:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(422, 'Expanded spreadsheet exceeds 30 MB')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book.active
                if sheet.max_column > 50 or sheet.max_row > MAX_ROWS + 1:
                    raise HTTPException(422, 'Maximum 50 columns and 2,000 data rows')
                rows = []
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        if cell.data_type == 'f':
                            raise HTTPException(422, 'Formula cells are not supported; upload values only')
                        values.append('' if cell.value is None else str(cell.value))
                    rows.append(values)
            finally:
                book.close()
        if len(rows) < 2:
            raise HTTPException(422, 'The file needs a header and at least one data row')
        headers = [str(h).strip() for h in rows[0]]
        if len(headers) > 50 or any(not h or len(h) > 100 for h in headers) or len(headers) != len(set(headers)):
            raise HTTPException(422, 'Use 1 to 50 unique, nonempty column headings')
        if any(len(row) > len(headers) for row in rows[1:]):
            raise HTTPException(422, 'A row has more values than column headings')
        return headers, [dict(zip(headers, row)) for row in rows[1:]], suffix[1:]
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Unable to read file. Use an unencrypted UTF-8 CSV or valid XLSX workbook')


def validate_rows(db, org, raw_rows, mapping):
    if not isinstance(mapping, dict) or any(k not in FIELDS or not isinstance(v, str) for k, v in mapping.items()) or 'full_name' not in mapping:
        raise HTTPException(422, 'Map at least full_name using supported fields')
    results, seen_email, seen_phone = [], set(), set()
    for number, row in enumerate(raw_rows, 2):
        values = {field: str(row.get(header, '')).strip() for field, header in mapping.items()}
        try:
            data = LeadInput(**values)
            match = duplicate(db, org, data)
            repeated = bool((data.email and data.email in seen_email) or (data.phone and data.phone in seen_phone))
            status = 'duplicate' if match or repeated else 'valid'
            if data.email:
                seen_email.add(data.email)
            if data.phone:
                seen_phone.add(data.phone)
            results.append({'row_number': number, 'status': status, 'data': data.model_dump(mode='json'), 'message': 'Matching email or phone; skipped on import' if status == 'duplicate' else ''})
        except ValidationError as error:
            message = '; '.join(f"{e['loc'][0]}: {e['msg']}" for e in error.errors(include_input=False, include_url=False))
            results.append({'row_number': number, 'status': 'failed', 'data': {}, 'message': message[:500]})
    return results


def preview_job(db, ctx, filename, content, mapping):
    if not isinstance(mapping, dict) or any(k not in FIELDS or not isinstance(v, str) for k, v in mapping.items()) or 'full_name' not in mapping:
        raise HTTPException(422, 'Map at least full_name using supported fields')
    headers, raw_rows, source = parse_file(filename, content)
    if any(header not in headers for header in mapping.values()):
        raise HTTPException(422, 'Mapped column is missing from this file')
    rows = validate_rows(db, ctx['organization_id'], raw_rows, mapping)
    counts = {kind: sum(r['status'] == kind for r in rows) for kind in ['valid', 'duplicate', 'failed']}
    job = ImportJob(organization_id=ctx['organization_id'], uploaded_by=ctx['user_id'], filename=Path(filename).name[:240], rows={'source': source, 'items': rows}, counts=counts)
    db['import_jobs'].insert_one(job)

    error_docs = [
        ImportErrorRow(organization_id=ctx['organization_id'], import_job_id=job.id, row_number=row['row_number'], error_code=row['status'], message=row['message'])
        for row in rows if row['status'] != 'valid'
    ]
    if error_docs:
        db['import_error_rows'].insert_many(error_docs)

    return job, rows


def commit_job(db, ctx, job):
    if job.status == 'completed':
        return job
    imported, duplicates = 0, job.counts.get('duplicate', 0)
    job_items = job.rows.get('items', []) if isinstance(job.rows, dict) else []
    error_docs = []

    for row in job_items:
        if row.get('status') != 'valid':
            continue
        data = LeadInput(**row.get('data', {}))
        if duplicate(db, ctx['organization_id'], data):
            duplicates += 1
            error_docs.append(ImportErrorRow(
                organization_id=ctx['organization_id'],
                import_job_id=job.id,
                row_number=row['row_number'],
                error_code='duplicate',
                message='Contact was added after preview; skipped'
            ))
            continue
        create_lead(db, ctx['organization_id'], data, job.rows.get('source', 'csv'), ctx['user_id'])
        imported += 1

    if error_docs:
        db['import_error_rows'].insert_many(error_docs)

    new_counts = {**job.counts, 'imported': imported, 'duplicate': duplicates}
    db['import_jobs'].update_one(
        {'$or': [{'_id': job.id}, {'id': job.id}]},
        {'$set': {'counts': new_counts, 'status': 'completed', 'rows': []}}
    )
    job.counts = new_counts
    job.status = 'completed'
    job.rows = []
    return job


def safe_csv_cell(value):
    value = str(value)
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else value
